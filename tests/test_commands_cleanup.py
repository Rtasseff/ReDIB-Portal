"""
commands-cleanup: management commands that are safe to use and honest when
they fail (backlog #82, #83, #88, #91).

- populate_redib_users creates users with no usable password (#82), has no
  --sync (#83), validates role names, ORCID and phone before writing, and
  loads all-or-nothing (#88 c/d).
- populate_redib_equipment loads all-or-nothing and leaves technical_specs
  alone when the file has no such column (#88 c/e).
- setup_base_database exits non-zero on a failed step (#88 a).
- export_redib_users writes the users.tsv format back out of the database,
  with retired roles in their own column (#91, fixes #81).
- send_test_emails renders the evaluator digests and --cleanup removes only
  the evaluator role it granted (#88 f).
"""
import csv
import io
import os
import tempfile
from pathlib import Path
from unittest import mock

from allauth.account.models import EmailAddress
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core import mail
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse

from core.management.commands.export_redib_users import COLUMNS
from core.models import Equipment, Node, Organization, UserRole

User = get_user_model()

USERS_HEADER = [
    'email', 'first_name', 'last_name', 'organization_name', 'orcid', 'phone',
    'position', 'is_staff', 'is_active', 'roles', 'areas', 'auto_data_consent',
]
EQUIPMENT_HEADER = ['node_code', 'name', 'category', 'description', 'area', 'is_essential', 'is_active']


def _write_tsv(header, rows):
    """Temp TSV with CRLF endings, like the real data/ files. Rows are dicts."""
    handle = tempfile.NamedTemporaryFile(
        mode='w', suffix='.tsv', delete=False, encoding='utf-8', newline=''
    )
    writer = csv.DictWriter(handle, fieldnames=header, delimiter='\t', lineterminator='\r\n')
    writer.writeheader()
    for row in rows:
        writer.writerow({k: row.get(k, '') for k in header})
    handle.close()
    return handle.name


def _user_row(email, **fields):
    row = {'email': email, 'first_name': 'First', 'last_name': 'Last',
           'is_active': 'TRUE', 'roles': 'evaluator', 'areas': 'clinical'}
    row.update(fields)
    return row


class _LoaderMixin:
    def setUp(self):
        self.org = Organization.objects.create(name='Test Org', short_name='TO')
        self.node = Node.objects.create(code='TEST-NODE', organization=self.org, location='Testville')

    def load_users(self, *rows, **options):
        path = _write_tsv(USERS_HEADER, rows)
        try:
            call_command('populate_redib_users', tsv=path, stdout=io.StringIO(), **options)
        finally:
            os.unlink(path)


class NoDefaultPasswordTest(_LoaderMixin, TestCase):
    """#82: a loaded user has no usable password and resets their own."""

    def test_new_user_has_unusable_password(self):
        self.load_users(_user_row('new@example.org'))
        user = User.objects.get(email='new@example.org')
        self.assertFalse(user.has_usable_password())
        self.assertFalse(user.check_password('changeme123'))

    def test_existing_user_keeps_their_password(self):
        User.objects.create_user(email='old@example.org', password='their-own')
        self.load_users(_user_row('old@example.org'))
        self.assertTrue(User.objects.get(email='old@example.org').check_password('their-own'))

    def test_output_never_prints_a_password(self):
        out = io.StringIO()
        path = _write_tsv(USERS_HEADER, [_user_row('new@example.org')])
        try:
            call_command('populate_redib_users', tsv=path, stdout=out)
        finally:
            os.unlink(path)
        self.assertNotIn('changeme', out.getvalue())
        self.assertIn('Forgot password', out.getvalue())

    def test_forgot_password_mails_a_loaded_user(self):
        """allauth looks users up by email and is_active only, so the reset
        form works for an unusable password."""
        self.load_users(_user_row('new@example.org'))
        self.assertTrue(EmailAddress.objects.get(email='new@example.org').verified)

        response = self.client.post(reverse('account_reset_password'), {'email': 'new@example.org'})

        self.assertEqual(response.status_code, 302)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ['new@example.org'])
        self.assertIn('/accounts/password/reset/key/', mail.outbox[0].body)


class NoSyncForUsersTest(_LoaderMixin, TestCase):
    """#83: the users loader has no --sync; the others keep theirs."""

    def test_users_loader_rejects_sync(self):
        path = _write_tsv(USERS_HEADER, [_user_row('a@example.org')])
        try:
            with self.assertRaises(CommandError):
                call_command('populate_redib_users', '--sync', tsv=path, stdout=io.StringIO())
        finally:
            os.unlink(path)

    def test_user_not_in_file_stays_active(self):
        applicant = User.objects.create_user(email='applicant@example.org', password='x')
        self.load_users(_user_row('a@example.org'))
        applicant.refresh_from_db()
        self.assertTrue(applicant.is_active)

    def test_equipment_loader_keeps_sync(self):
        path = _write_tsv(EQUIPMENT_HEADER, [])
        try:
            call_command('populate_redib_equipment', '--sync', tsv=path, stdout=io.StringIO())
        finally:
            os.unlink(path)


class UserLoaderValidationTest(_LoaderMixin, TestCase):
    """#88 c/d: nothing is written unless every row is good."""

    def assert_load_fails(self, *rows, message):
        before = (User.objects.count(), UserRole.objects.count())
        with self.assertRaisesMessage(CommandError, message):
            self.load_users(*rows)
        self.assertEqual((User.objects.count(), UserRole.objects.count()), before)

    def test_unknown_role_aborts_and_lists_valid_names(self):
        self.assert_load_fails(
            _user_row('a@example.org'),
            _user_row('b@example.org', roles='evalutor'),
            message="unknown role 'evalutor'. Valid roles: applicant, node_coordinator, evaluator",
        )

    def test_invalid_orcid_aborts(self):
        self.assert_load_fails(
            _user_row('a@example.org', orcid='0000-0002-1234'),
            message='invalid orcid',
        )

    def test_invalid_phone_aborts(self):
        self.assert_load_fails(
            _user_row('a@example.org', phone='call me maybe'),
            message='invalid phone',
        )

    def test_valid_orcid_and_phone_load(self):
        self.load_users(_user_row('a@example.org', orcid='0000-0002-1234-567X', phone='+34 91 0672 228'))
        self.assertEqual(User.objects.get(email='a@example.org').orcid, '0000-0002-1234-567X')

    def test_bad_row_after_good_rows_writes_nothing(self):
        """A failure only found while writing (an unknown organization) rolls
        back the rows already written above it."""
        self.assert_load_fails(
            _user_row('a@example.org'),
            _user_row('b@example.org', roles=f'node_coordinator:{self.node.code}', areas=''),
            _user_row('c@example.org', organization_name='No Such Org'),
            message='No Such Org',
        )
        self.assertFalse(User.objects.filter(email='a@example.org').exists())

    def test_dry_run_writes_nothing(self):
        self.load_users(_user_row('a@example.org'), dry_run=True)
        self.assertFalse(User.objects.exists())


class EquipmentLoaderTest(_LoaderMixin, TestCase):
    """#88 c/e."""

    def load_equipment(self, rows, header=EQUIPMENT_HEADER):
        path = _write_tsv(header, rows)
        try:
            call_command('populate_redib_equipment', tsv=path, stdout=io.StringIO())
        finally:
            os.unlink(path)

    def _row(self, name, node_code=None, **fields):
        row = {'node_code': node_code or self.node.code, 'name': name, 'category': 'mri',
               'is_essential': 'TRUE', 'is_active': 'TRUE'}
        row.update(fields)
        return row

    def test_technical_specs_kept_without_column(self):
        Equipment.objects.create(node=self.node, name='MRI 7T', category='mri', technical_specs='7 tesla, 30 cm bore')
        self.load_equipment([self._row('MRI 7T', description='Updated')])
        equipment = Equipment.objects.get(name='MRI 7T')
        self.assertEqual(equipment.description, 'Updated')
        self.assertEqual(equipment.technical_specs, '7 tesla, 30 cm bore')

    def test_technical_specs_written_with_column(self):
        Equipment.objects.create(node=self.node, name='MRI 7T', category='mri', technical_specs='old')
        self.load_equipment(
            [self._row('MRI 7T', technical_specs='new')],
            header=EQUIPMENT_HEADER + ['technical_specs'],
        )
        self.assertEqual(Equipment.objects.get(name='MRI 7T').technical_specs, 'new')

    def test_unknown_node_writes_nothing(self):
        with self.assertRaisesMessage(CommandError, 'NO-SUCH-NODE'):
            self.load_equipment([self._row('MRI 7T'), self._row('PET', node_code='NO-SUCH-NODE')])
        self.assertFalse(Equipment.objects.exists())


class SetupBaseDatabaseFailureTest(TestCase):
    """#88 a: a failed step is a non-zero exit, not a printed 'Failed:'."""

    def test_failed_step_raises(self):
        real_call_command = call_command

        def fail_on_users(name, *args, **kwargs):
            if name == 'populate_redib_users':
                raise CommandError('bad users.tsv')
            return real_call_command(name, *args, stdout=io.StringIO(), **kwargs)

        with mock.patch(
            'core.management.commands.setup_base_database.call_command', side_effect=fail_on_users
        ):
            with self.assertRaisesMessage(CommandError, 'Step 3 failed: bad users.tsv'):
                call_command('setup_base_database', stdout=io.StringIO())


class ExportUsersTest(TestCase):
    """#91: export_redib_users writes the users.tsv format from the DB."""

    def export(self, *args):
        out = io.StringIO()
        call_command('export_redib_users', *args, stdout=out)
        return out.getvalue()

    def rows(self, text):
        return list(csv.DictReader(io.StringIO(text, newline=''), delimiter='\t'))

    def load_reference_data(self):
        for command in ('populate_redib_organizations', 'populate_redib_nodes', 'populate_redib_users'):
            call_command(command, stdout=io.StringIO())

    def test_round_trip_of_committed_users_tsv(self):
        """Load data/users.tsv, export, and every value matches as the loader
        reads it. Two things differ in the text: blank booleans come back as
        FALSE, and a row with no role that isn't staff isn't exported."""
        self.load_reference_data()
        text = self.export()

        with open(Path(settings.BASE_DIR) / 'data' / 'users.tsv', encoding='utf-8', newline='') as f:
            committed = list(csv.DictReader(f, delimiter='\t'))
        exported = {r['email']: r for r in self.rows(text)}

        def as_loaded(row):
            values = {}
            for column in USERS_HEADER:
                value = (row.get(column) or '').strip()
                if column in ('is_staff', 'is_active', 'auto_data_consent'):
                    value = value.upper() in ('TRUE', '1', 'YES')
                values[column] = value
            values['email'] = values['email'].lower()
            return values

        expected = [
            as_loaded(r) for r in committed
            if r['roles'].strip() or as_loaded(r)['is_staff']
        ]
        self.assertEqual(sorted(exported), sorted(r['email'] for r in expected))
        for row in expected:
            self.assertEqual(as_loaded(exported[row['email']]), row, row['email'])
            self.assertEqual(exported[row['email']]['retired_roles'], '')

    def test_format(self):
        self.load_reference_data()
        text = self.export()
        lines = text.split('\r\n')
        self.assertEqual(lines[0].split('\t'), COLUMNS)
        self.assertEqual(COLUMNS[:-1], USERS_HEADER)
        self.assertEqual(lines[-1], '')  # ends with CRLF
        self.assertNotIn('\n', text.replace('\r\n', ''))
        self.assertFalse(text.startswith('﻿'))
        emails = [r['email'] for r in self.rows(text)]
        self.assertEqual(emails, sorted(emails))
        for row in self.rows(text):
            for column in ('is_staff', 'is_active', 'auto_data_consent'):
                self.assertIn(row[column], ('TRUE', 'FALSE'))

    def test_output_file_is_utf8_without_bom(self):
        self.load_reference_data()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'users.tsv')
            call_command('export_redib_users', output=path, stdout=io.StringIO(), stderr=io.StringIO())
            data = Path(path).read_bytes()
        self.assertFalse(data.startswith(b'\xef\xbb\xbf'))
        self.assertIn('Ángel Manuel'.encode('utf-8'), data)
        self.assertIn(b'\r\n', data)

    def test_xlsx(self):
        from openpyxl import load_workbook

        self.load_reference_data()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'users.xlsx')
            call_command('export_redib_users', output=path, format='xlsx',
                         stdout=io.StringIO(), stderr=io.StringIO())
            sheet = load_workbook(path).active
            header = [c.value for c in sheet[1]]
            n_rows = sheet.max_row - 1
        self.assertEqual(header, COLUMNS)
        self.assertEqual(n_rows, len(self.rows(self.export())))

    def test_who_is_exported(self):
        org = Organization.objects.create(name='Org', short_name='O')
        node = Node.objects.create(code='N1', organization=org, location='X')
        applicant = User.objects.create_user(email='applicant@example.org')
        UserRole.objects.create(user=applicant, role='applicant')
        User.objects.create_user(email='nobody@example.org')
        User.objects.create_user(email='staff@example.org', is_staff=True)
        both = User.objects.create_user(email='both@example.org')
        UserRole.objects.create(user=both, role='applicant')
        UserRole.objects.create(user=both, role='node_coordinator', node=node)
        UserRole.objects.create(user=both, role='coordinator')

        rows = {r['email']: r for r in self.rows(self.export())}

        self.assertEqual(sorted(rows), ['both@example.org', 'staff@example.org'])
        self.assertEqual(rows['both@example.org']['roles'], 'node_coordinator:N1;coordinator')
        self.assertEqual(rows['staff@example.org']['roles'], '')

    def test_retired_role_is_recorded_not_regranted(self):
        """#81: a retired evaluator is exported under retired_roles, and
        loading the export back leaves the role inactive. The loader reads
        with DictReader, so it ignores the extra column."""
        user = User.objects.create_user(email='former@example.org', first_name='Former', last_name='Evaluator')
        role = UserRole.objects.create(user=user, role='evaluator', areas='clinical;preclinical', is_active=False)

        text = self.export()
        row = self.rows(text)[0]
        self.assertEqual(row['roles'], '')
        self.assertEqual(row['retired_roles'], 'evaluator')
        self.assertEqual(row['areas'], 'clinical;preclinical')

        with tempfile.NamedTemporaryFile('w', suffix='.tsv', delete=False, encoding='utf-8', newline='') as f:
            f.write(text)
        try:
            call_command('populate_redib_users', tsv=f.name, stdout=io.StringIO())
        finally:
            os.unlink(f.name)
        role.refresh_from_db()
        self.assertFalse(role.is_active)
        self.assertEqual(UserRole.objects.filter(user=user).count(), 1)


class SendTestEmailsTest(TestCase):
    """#88 f."""

    def setUp(self):
        call_command('seed_email_templates', stdout=io.StringIO())
        org = Organization.objects.create(name='Org', short_name='O')
        Node.objects.create(code='N1', organization=org, location='X')
        self.user = User.objects.create_user(email='you@example.org', first_name='You', last_name='Tester')

    def run_command(self, *args):
        call_command('send_test_emails', *args, to='you@example.org', stdout=io.StringIO())

    def test_digest_subjects_render(self):
        self.run_command()
        subjects = [m.subject for m in mail.outbox]
        self.assertIn('ReDIB COA: Evaluation Reminder (1 Pending)', subjects)
        self.assertIn('ReDIB COA: 1 Evaluation Overdue', subjects)
        reminder = next(m for m in mail.outbox if 'Reminder (1 Pending)' in m.subject)
        self.assertIn('3 days remaining', reminder.body)

    def test_cleanup_removes_the_role_it_granted(self):
        self.run_command()
        self.assertTrue(UserRole.objects.filter(user=self.user, role='evaluator').exists())
        self.run_command('--cleanup')
        self.assertFalse(UserRole.objects.filter(user=self.user).exists())

    def test_cleanup_keeps_a_role_the_person_already_had(self):
        role = UserRole.objects.create(user=self.user, role='evaluator', areas='clinical')
        self.run_command()
        self.run_command('--cleanup')
        role.refresh_from_db()
        self.assertEqual(role.areas, 'clinical')
        self.assertTrue(role.is_active)
